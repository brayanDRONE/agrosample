"""Lectura y normalizacion de archivos Excel para validacion de SDP."""

import re
import unicodedata
from datetime import date, datetime

from openpyxl import load_workbook
import openpyxl.reader.excel
from openpyxl.styles.colors import RGB

# Patch openpyxl para tolerar colores aRGB malformados o stylesheets invalidos en archivos de terceros (ej. SIF / SAG)
_orig_rgb_set = RGB.__set__
def _safe_rgb_set(self, instance, value):
    if value is not None and isinstance(value, str):
        if len(value) == 6:
            value = '00' + value
        elif len(value) != 8:
            value = '00000000'
    try:
        super(RGB, self).__set__(instance, value)
    except Exception:
        pass
RGB.__set__ = _safe_rgb_set

_orig_apply_stylesheet = openpyxl.reader.excel.apply_stylesheet
def _safe_apply_stylesheet(archive, wb):
    try:
        _orig_apply_stylesheet(archive, wb)
    except Exception:
        pass
openpyxl.reader.excel.apply_stylesheet = _safe_apply_stylesheet


FIELD_ALIASES = {
    'csg': ('csg', 'codigo csg', 'codigo sag', 'codigosag', 'codigo'),
    'provincia': ('provincia origen', 'provincia', 'provincia pack'),
    'comuna': ('comuna origen', 'comuna', 'comuna pack'),
    'variedad_rotulada': (
        'variedad rotulada', 'variedadrotulada', 'variedad comercial', 'variedad',
        'variedad real', 'variedadreal', 'variedad agronomica', 'variedad agron mica'
    ),
    'sdp': ('sdp', 'sdp sector', 'sitio de produccion', 'sitio produccion'),
}

OPTIONAL_ALIASES = {
    'cajas': (
        'cajas', 'caja', 'cant cajas', 'cantidad cajas', 'cant. cajas', 'cant de cajas',
        'cantidad de cajas', 'tot cajas', 'total cajas', 'nro cajas', 'num cajas', 'bultos',
        'cajas embaladas', 'cajas inspeccion', 'cajas reales'
    ),
    'fecha': (
        'fec pack', 'fec packing', 'fecha pack', 'fecha packing', 'fecpack',
        'fecha', 'fecha embalaje', 'fec embalaje', 'fecha de embalaje',
        'fecha proceso', 'fec proceso', 'fecha cosecha', 'fec cosecha',
        'fecha produccion', 'fec produccion', 'fecha elaboracion', 'fec elaboracion',
        'fec inspec', 'fec inspeccion', 'fecha inspeccion', 'fec mov', 'fecha mov'
    ),
}


def normalize_header(value):
    """Normaliza un encabezado para compararlo sin acentos ni formato."""
    text = '' if value is None else str(value)
    text = unicodedata.normalize('NFKD', text)
    text = ''.join(char for char in text if not unicodedata.combining(char))
    return re.sub(r'[^a-z0-9]+', ' ', text.lower()).strip()


def normalize_value(value):
    """Convierte valores Excel a texto estable sin perder codigos SAG."""
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalize_date_value(value):
    """Normaliza un valor de fecha a formato DD/MM/YYYY."""
    if value is None:
        return ''
    if isinstance(value, (datetime, date)):
        return value.strftime('%d/%m/%Y')
    val_str = str(value).strip()
    match_iso = re.match(r'^(\d{4})[-/](\d{1,2})[-/](\d{1,2})', val_str)
    if match_iso:
        year, month, day = match_iso.groups()
        return f'{int(day):02d}/{int(month):02d}/{year}'
    match_latam = re.match(r'^(\d{1,2})[-/](\d{1,2})[-/](\d{2,4})', val_str)
    if match_latam:
        day, month, year = match_latam.groups()
        if len(year) == 2:
            year = f'20{year}'
        return f'{int(day):02d}/{int(month):02d}/{year}'
    return val_str


def normalize_cajas_value(value):
    """Normaliza un valor de cajas a entero."""
    if value is None:
        return 0
    if isinstance(value, (int, float)):
        return int(round(value))
    val_str = str(value).strip()
    digits = re.search(r'\d+', val_str.replace('.', '').replace(',', ''))
    if digits:
        return int(digits.group(0))
    return 0


def extract_workbook_metadata(worksheet):
    """Extrae el numero de lote desde el bloque superior de la hoja."""
    metadata = {'numero_lote': ''}
    top_rows = list(worksheet.iter_rows(min_row=1, max_row=6, values_only=True))

    for row in top_rows:
        for index, value in enumerate(row):
            label = normalize_header(value)
            if 'inspeccion sag' not in label:
                continue
            for candidate in row[index + 1:]:
                normalized = normalize_value(candidate)
                if normalized:
                    metadata['numero_lote'] = normalized
                    return metadata

    return metadata


def resolve_columns(headers):
    normalized = {normalize_header(header): index for index, header in enumerate(headers)}
    columns = {}
    missing = []

    for field, aliases in FIELD_ALIASES.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match is None:
            missing.append(field)
            continue
        columns[field] = match

    optional_columns = {}
    for field, aliases in OPTIONAL_ALIASES.items():
        match = next((normalized[alias] for alias in aliases if alias in normalized), None)
        if match is None and field == 'fecha':
            # Fallback: buscar encabezados que empiecen con 'fec' o contengan 'fecha'
            for h, idx in normalized.items():
                if h.startswith('fec') or 'fecha' in h:
                    match = idx
                    break
        if match is not None:
            optional_columns[field] = match

    return columns, missing, optional_columns


def parse_sdp_workbook(file_object):
    """Lee todas las hojas y devuelve filas normalizadas y SDP unicos."""
    workbook = load_workbook(file_object, read_only=True, data_only=True)
    rows = []
    sheets = []
    errors = []
    metadata = {'numero_lote': ''}

    try:
        for worksheet in workbook.worksheets:
            sheets.append(worksheet.title)
            if not metadata['numero_lote']:
                metadata = extract_workbook_metadata(worksheet)
            values = worksheet.iter_rows(values_only=True)
            header_row_number = None
            headers = None
            buffered_rows = []
            for row_number, candidate in enumerate(values, start=1):
                buffered_rows.append((row_number, candidate))
                columns, missing, _ = resolve_columns(candidate)
                if len(missing) <= 1 and len(columns) >= 5:
                    headers = candidate
                    header_row_number = row_number
                    break
                if row_number >= 30:
                    break

            if headers is None:
                continue

            columns, missing, optional_columns = resolve_columns(headers)
            if missing:
                errors.append(
                    f"Hoja '{worksheet.title}': faltan columnas requeridas: {', '.join(missing)}"
                )
                continue

            for row_number, values_row in enumerate(values, start=header_row_number + 1):
                record = {
                    field: normalize_value(values_row[index] if index < len(values_row) else None)
                    for field, index in columns.items()
                }
                if not any(record.values()):
                    continue

                if 'fecha' in optional_columns:
                    f_idx = optional_columns['fecha']
                    val = values_row[f_idx] if f_idx < len(values_row) else None
                    record['fecha'] = normalize_date_value(val)
                else:
                    record['fecha'] = ''

                if 'cajas' in optional_columns:
                    c_idx = optional_columns['cajas']
                    val = values_row[c_idx] if c_idx < len(values_row) else None
                    record['cajas'] = normalize_cajas_value(val)
                else:
                    record['cajas'] = 0

                record['hoja'] = worksheet.title
                record['fila_excel'] = row_number
                rows.append(record)
    finally:
        workbook.close()

    unique_sdps = list(dict.fromkeys(row['sdp'] for row in rows if row['sdp']))
    return {
        'sheets': sheets,
        'metadata': metadata,
        'rows': rows,
        'unique_sdps': unique_sdps,
        'errors': errors,
    }